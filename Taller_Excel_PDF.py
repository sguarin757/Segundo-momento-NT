from openpyxl import Workbook, load_workbook
from tabulate import tabulate
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas

# 1. CREAR EL EXCEL
ventas = [
    ("Laptop", "Computadores", 5, 1200),
    ("Mouse", "Accesorios", 40, 25),
    ("Teclado", "Accesorios", 25, 45),
    ("Monitor", "Pantallas", 10, 350),
    ("Impresora", "Oficina", 8, 180),
    ("Auriculares", "Accesorios", 30, 85),
    ("Disco Duro", "Almacenamiento", 15, 60),
    ("Camara Web", "Accesorios", 22, 50),
    ("Silla", "Oficina", 6, 220),
    ("Escritorio", "Oficina", 4, 400),
]

wb = Workbook()
ws = wb.active
ws.title = "Ventas"

ws.cell(row=1, column=1, value="Producto")
ws.cell(row=1, column=2, value="Categoria")
ws.cell(row=1, column=3, value="Cantidad")
ws.cell(row=1, column=4, value="Precio")
ws.cell(row=1, column=5, value="Total")

for fila, (producto, categoria, cantidad, precio) in enumerate(ventas, start=2):
    ws.cell(row=fila, column=1, value=producto)
    ws.cell(row=fila, column=2, value=categoria)
    ws.cell(row=fila, column=3, value=cantidad)
    ws.cell(row=fila, column=4, value=precio)
    ws.cell(row=fila, column=5, value=cantidad * precio)

wb.save("ventas.xlsx")
print("Excel creado: ventas.xlsx")

# 2. LEER EL EXCEL
wb = load_workbook("ventas.xlsx")
ws = wb.active

tabla = []
for fila in range(2, ws.max_row + 1):
    producto = ws.cell(row=fila, column=1).value
    categoria = ws.cell(row=fila, column=2).value
    cantidad = ws.cell(row=fila, column=3).value
    precio = ws.cell(row=fila, column=4).value
    total = ws.cell(row=fila, column=5).value
    tabla.append([producto, categoria, cantidad, precio, total])

# 3. CALCULOS
total_ventas = 0
total_unidades = 0
for fila in tabla:
    total_ventas = total_ventas + fila[4]
    total_unidades = total_unidades + fila[2]

promedio = total_ventas / total_unidades

# producto con mayor y menor total
mejor = tabla[0]
peor = tabla[0]
for fila in tabla:
    if fila[4] > mejor[4]:
        mejor = fila
    if fila[4] < peor[4]:
        peor = fila

# total por categoria
categorias = []
totales_cat = []
for fila in tabla:
    if fila[1] in categorias:
        posicion = categorias.index(fila[1])
        totales_cat[posicion] = totales_cat[posicion] + fila[4]
    else:
        categorias.append(fila[1])
        totales_cat.append(fila[4])

# 4. MOSTRAR EN CONSOLA CON TABULATE
print("\n=== DETALLE DE VENTAS ===")
print(tabulate(tabla, headers=["Producto", "Categoria", "Cantidad", "Precio", "Total"], tablefmt="grid"))

resumen = [
    ["Total de ventas", total_ventas],
    ["Unidades vendidas", total_unidades],
    ["Precio promedio por unidad", round(promedio, 2)],
    ["Producto con mas ingresos", mejor[0]],
    ["Producto con menos ingresos", peor[0]],
]
print("\n=== INDICADORES ===")
print(tabulate(resumen, tablefmt="grid"))

tabla_cat = []
for i in range(len(categorias)):
    tabla_cat.append([categorias[i], totales_cat[i]])
print("\n=== VENTAS POR CATEGORIA ===")
print(tabulate(tabla_cat, headers=["Categoria", "Total"], tablefmt="grid"))

# 5. GENERAR EL PDF AUTOMATICAMENTE
pdf = canvas.Canvas("reporte_ejecutivo_ventas.pdf", pagesize=letter)
y = 750 

pdf.setFont("Helvetica-Bold", 18)
pdf.drawString(50, y, "Reporte Ejecutivo de Ventas")
y = y - 40

pdf.setFont("Helvetica-Bold", 13)
pdf.drawString(50, y, "Indicadores")
y = y - 22
pdf.setFont("Helvetica", 11)
for item in resumen:
    pdf.drawString(60, y, item[0] + ": " + str(item[1]))
    y = y - 18

y = y - 15
pdf.setFont("Helvetica-Bold", 13)
pdf.drawString(50, y, "Ventas por categoria")
y = y - 22
pdf.setFont("Helvetica", 11)
for fila in tabla_cat:
    pdf.drawString(60, y, fila[0] + ": " + str(fila[1]))
    y = y - 18

y = y - 15
pdf.setFont("Helvetica-Bold", 13)
pdf.drawString(50, y, "Detalle de ventas")
y = y - 22

pdf.setFont("Helvetica-Bold", 10)
pdf.drawString(60, y, "Producto")
pdf.drawString(180, y, "Categoria")
pdf.drawString(300, y, "Cantidad")
pdf.drawString(370, y, "Precio")
pdf.drawString(440, y, "Total")
y = y - 16

pdf.setFont("Helvetica", 10)
for fila in tabla:
    pdf.drawString(60, y, str(fila[0]))
    pdf.drawString(180, y, str(fila[1]))
    pdf.drawString(300, y, str(fila[2]))
    pdf.drawString(370, y, str(fila[3]))
    pdf.drawString(440, y, str(fila[4]))
    y = y - 16

pdf.save()
print("\nPDF generado: reporte_ejecutivo_ventas.pdf")